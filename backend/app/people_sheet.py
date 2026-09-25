"""The people sheet: one workbook that bulk-creates per-person accounts and their access.

Tabs:
  People               one row per person: ID, name, department, designation, official and
                       personal email, status, then one column per service holding a role.
  Department defaults  department -> role per service, used where a person's cell is blank.
  Instructions         how to fill it in.
  Lists                hidden; feeds the dropdowns (departments, statuses, each service's roles).

Precedence for each person and service (owner decision 25 Sep 2026, D-2026-09-25-2):
  a filled-in cell  >  the department default  >  the service's lowest role.
A filled-in cell replaces an existing role; a default or lowest role only fills a gap; `none`
removes access; a blank cell never removes anything. The actual assignment runs through
roles_io.assign_roles, the same engine role files use.

Everything is computed against real rows and, for a dry run, rolled back by the caller, so the
preview is exactly what the apply will do. The apply re-reads the uploaded file rather than
trusting the preview, and skips only the (person, service) pairs the admin unticked.
"""
from __future__ import annotations

import io
import secrets
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timezone
from typing import BinaryIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import func, select
from sqlalchemy.orm import Session as OrmSession

from . import roles_io
from .departments import (
    CANONICAL, UNASSIGNED, canonical_department, clean_email, is_company_email, normalize_code,
)
from .models import Employee, Grant, Revocation, Service, ServiceRole, Session, User
from .onboarding import PersonalEmail, personal_email_for
from .provision import FUNCTIONAL_JOB_TITLE
from .security import hash_pin

PEOPLE_SHEET = "People"
DEFAULTS_SHEET = "Department defaults"
FIXED = ["Employee ID", "Full name", "Department", "Designation", "Official email", "Personal email", "Status"]
NOTES_HEADER = "Review notes"
STATUSES = ("active", "exited")
NONE = "none"
MAX_ROWS = 1000
PETROL = "005D7F"
PROPOSED_FILL = PatternFill("solid", fgColor="FFF4C2")


@dataclass
class ServiceCol:
    slug: str
    name: str
    roles: list[str]          # lowest access first
    lowest: str | None        # what an unconfigured person gets; None = nothing


def service_columns(db: OrmSession) -> list[ServiceCol]:
    """Every active service MM OS signs people into. External services (ERPNext, Twenty)
    keep their own logins, so the sheet has no column for them."""
    cols = []
    services = db.scalars(
        select(Service).where(Service.is_active.is_(True), Service.launch_mode != "external")
        .order_by(Service.sort_order, Service.name)
    ).all()
    for s in services:
        roles = list(s.roles)  # already ordered lowest access first (sort_order)
        if not roles:
            continue
        cols.append(ServiceCol(s.slug, s.name, [r.key for r in roles], next((r.key for r in roles if r.is_default), None)))
    return cols


def seed_service_columns() -> list[ServiceCol]:
    """The same columns built from the committed definitions (seed + role files), for making
    a pre-filled sheet on a machine with no MM OS database."""
    from .seed import LOWEST_ROLES, SERVICES
    cols = []
    for spec in SERVICES:
        if spec.get("launch_mode") == "external" or spec.get("is_active") is False:
            continue
        doc = roles_io.committed(spec["slug"])
        if doc:
            keys = [r["key"] for r in doc["roles"]]
            lowest = next((r["key"] for r in doc["roles"] if r.get("default")), None)
        else:
            keys = [k for k, _n, _d in spec["roles"]]
            lowest = LOWEST_ROLES.get(spec["slug"])
        if keys:
            cols.append(ServiceCol(spec["slug"], spec["name"], keys, lowest))
    return cols


# ── template ──────────────────────────────────────────────────────────────────
def _col_letter(i: int) -> str:
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def _style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF", name="Roboto")
        cell.fill = PatternFill("solid", fgColor=PETROL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"


INSTRUCTIONS = [
    ("MM OS people sheet", True),
    ("", False),
    ("One row per person on the People tab. Upload it at MM OS → Admin → People → Bulk upload. "
     "Nothing is written until you have reviewed the dry run and confirmed.", False),
    ("", False),
    ("Columns", True),
    ("Employee ID: MM followed by the number, e.g. MM115. No zeros in front, no dashes or spaces "
     "(the upload corrects these and lists every correction).", False),
    ("Full name, Department: required. Pick the department from the list.", False),
    ("Official email: the company address (@m-mines.com). The person signs in with Google on this "
     "address directly.", False),
    ("Personal email: optional Gmail. It only works after the person types their own employee "
     "code the first time they sign in with it.", False),
    ("Status: active (default) or exited. Exited switches the account off.", False),
    ("One column per service: pick a role from the list.", False),
    ("", False),
    ("What a service cell means", True),
    ("A role: the person gets exactly that role, replacing any role they have now.", False),
    ("Blank: the department's default from the Department defaults tab, but only if the person "
     "has no role on that service yet. With no department default: the service's lowest role.", False),
    ("none: remove the person's access to that service.", False),
    ("", False),
    ("First sign-in", True),
    ("Everyone signs in with Google, then confirms their employee code and chooses a 4–8 digit "
     "PIN once. After that, Google or employee code + PIN both work.", False),
    ("Someone who signs in with a company address but is not on this sheet gets department "
     "Unassigned and only each service's lowest role. Any other unknown address is refused.", False),
    ("", False),
    ("Review notes (last column) are for reviewers and are ignored on upload. Yellow cells are "
     "proposals to confirm.", False),
    ("Platform administrators cannot be made from this sheet.", False),
]


def build_workbook(
    services: list[ServiceCol],
    *,
    people: list[dict] | None = None,
    defaults: dict[str, dict[str, str]] | None = None,
    proposed: set[tuple[str, str]] | None = None,
) -> bytes:
    """The template, optionally pre-filled. `people` rows are dicts with the FIXED fields as
    keys plus `roles` ({slug: key}) and `notes`. `proposed` marks (department, slug) default
    cells to highlight for review."""
    wb = Workbook()
    ws_i = wb.active
    ws_i.title = "Instructions"
    ws_p = wb.create_sheet(PEOPLE_SHEET)
    ws_d = wb.create_sheet(DEFAULTS_SHEET)
    ws_l = wb.create_sheet("Lists")

    # Lists: A departments, B statuses, then one column per service.
    departments = list(CANONICAL)
    ws_l.cell(row=1, column=1, value="Departments")
    for i, d in enumerate(departments, start=2):
        ws_l.cell(row=i, column=1, value=d)
    ws_l.cell(row=1, column=2, value="Status")
    for i, s in enumerate(STATUSES, start=2):
        ws_l.cell(row=i, column=2, value=s)
    list_ranges: dict[str, str] = {}
    for j, svc in enumerate(services):
        col = 3 + j
        ws_l.cell(row=1, column=col, value=svc.name)
        options = svc.roles + [NONE]
        for i, key in enumerate(options, start=2):
            ws_l.cell(row=i, column=col, value=key)
        letter = _col_letter(col - 1)
        list_ranges[svc.slug] = f"Lists!${letter}$2:${letter}${len(options) + 1}"
    ws_l.sheet_state = "hidden"
    dept_range = f"Lists!$A$2:$A${len(departments) + 1}"
    status_range = f"Lists!$B$2:$B${len(STATUSES) + 1}"

    def _validate(ws, col_idx: int, formula: str, first_row: int = 2) -> None:
        dv = DataValidation(type="list", formula1=f"={formula}", allow_blank=True, showErrorMessage=True,
                            errorTitle="Pick from the list", error="Choose a value from the dropdown.")
        letter = _col_letter(col_idx)
        dv.add(f"{letter}{first_row}:{letter}{MAX_ROWS}")
        ws.add_data_validation(dv)

    # People
    headers = FIXED + [s.name for s in services] + [NOTES_HEADER]
    for c, h in enumerate(headers, start=1):
        ws_p.cell(row=1, column=c, value=h)
    _style_header(ws_p, len(headers))
    widths = [12, 26, 18, 26, 30, 30, 10] + [16] * len(services) + [60]
    for c, w in enumerate(widths):
        ws_p.column_dimensions[_col_letter(c)].width = w
    _validate(ws_p, FIXED.index("Department"), dept_range)
    _validate(ws_p, FIXED.index("Status"), status_range)
    for j, svc in enumerate(services):
        _validate(ws_p, len(FIXED) + j, list_ranges[svc.slug])
    for r, person in enumerate(people or [], start=2):
        for c, h in enumerate(FIXED, start=1):
            ws_p.cell(row=r, column=c, value=person.get(h) or None)
        for j, svc in enumerate(services):
            key = (person.get("roles") or {}).get(svc.slug)
            if key:
                ws_p.cell(row=r, column=len(FIXED) + 1 + j, value=key)
        notes = person.get("notes")
        if notes:
            cell = ws_p.cell(row=r, column=len(headers), value=notes)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    # Department defaults
    dheaders = ["Department"] + [s.name for s in services]
    for c, h in enumerate(dheaders, start=1):
        ws_d.cell(row=1, column=c, value=h)
    _style_header(ws_d, len(dheaders))
    ws_d.column_dimensions["A"].width = 20
    for j, svc in enumerate(services):
        ws_d.column_dimensions[_col_letter(1 + j)].width = 16
        _validate(ws_d, 1 + j, list_ranges[svc.slug])
    for r, dept in enumerate(departments, start=2):
        ws_d.cell(row=r, column=1, value=dept)
        for j, svc in enumerate(services):
            key = (defaults or {}).get(dept, {}).get(svc.slug)
            if key:
                cell = ws_d.cell(row=r, column=2 + j, value=key)
                if proposed and (dept, svc.slug) in proposed:
                    cell.fill = PROPOSED_FILL
    lowest_row = len(departments) + 3
    ws_d.cell(row=lowest_row, column=1, value="Lowest role (blank cell, no default)").font = Font(italic=True)
    for j, svc in enumerate(services):
        ws_d.cell(row=lowest_row, column=2 + j, value=svc.lowest or "nothing").font = Font(italic=True, color="666666")

    # Instructions
    ws_i.column_dimensions["A"].width = 110
    for r, (text, bold) in enumerate(INSTRUCTIONS, start=1):
        cell = ws_i.cell(row=r, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if bold:
            cell.font = Font(bold=True, size=13 if r == 1 else 11, color=PETROL)

    wb.active = 1  # open on People
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ── parsing and clean-up ──────────────────────────────────────────────────────
@dataclass
class SheetRow:
    row: int
    raw_code: str
    code: str | None = None
    name: str = ""
    department: str = UNASSIGNED
    raw_department: str = ""
    designation: str = ""
    official: str | None = None
    personal: str | None = None
    status: str = "active"
    roles: dict[str, str] = field(default_factory=dict)
    errors: list[str] = field(default_factory=list)
    fixes: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    duplicate_of: int | None = None

    def identity(self) -> tuple:
        return (self.name.lower(), self.department, self.designation.lower(), self.official,
                self.personal, self.status, tuple(sorted(self.roles.items())))


@dataclass
class ParsedSheet:
    rows: list[SheetRow]
    defaults: dict[str, dict[str, str]]
    problems: list[str]


def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())


def _header_map(headers: list[str], services: list[ServiceCol]) -> tuple[dict[str, int], dict[str, int], list[str]]:
    fixed, svc, ignored = {}, {}, []
    by_name = {s.name.lower(): s.slug for s in services} | {s.slug.lower(): s.slug for s in services}
    wanted = {h.lower(): h for h in FIXED}
    for i, raw in enumerate(headers):
        h = raw.strip().rstrip("*").strip()
        if not h:
            continue
        if h.lower() in wanted:
            fixed[wanted[h.lower()]] = i
        elif h.lower() in by_name:
            svc[by_name[h.lower()]] = i
        elif h.lower() != NOTES_HEADER.lower():
            ignored.append(h)
    return fixed, svc, ignored


def _role_key(value: str, svc: ServiceCol, names: dict[str, str]) -> str | None:
    v = value.strip().lower()
    if v in svc.roles or v == NONE:
        return v
    return names.get(v)


def parse_workbook(source: BinaryIO, services: list[ServiceCol], role_names: dict[str, dict[str, str]]) -> ParsedSheet:
    """Read and clean every row. Never raises for a bad row: the problem lands on the row."""
    wb = load_workbook(source, read_only=True, data_only=True)
    if PEOPLE_SHEET not in wb.sheetnames:
        raise ValueError(f"the workbook has no {PEOPLE_SHEET!r} tab -- download the template from MM OS")
    problems: list[str] = []
    by_slug = {s.slug: s for s in services}

    grid = list(wb[PEOPLE_SHEET].iter_rows(values_only=True))
    headers = [_s(h) for h in (grid[0] if grid else [])]
    fixed, svc_cols, ignored = _header_map(headers, services)
    missing = [h for h in ("Employee ID", "Full name", "Department") if h not in fixed]
    if missing:
        raise ValueError(f"the People tab is missing columns: {', '.join(missing)}")
    if ignored:
        problems.append(f"columns ignored (not a known field or service): {', '.join(ignored)}")
    for s in services:
        if s.slug not in svc_cols:
            problems.append(f"no column for {s.name}: nobody's role on it is set from this sheet, "
                            "people only get its department default or lowest role")

    def cell(values, name):
        i = fixed.get(name)
        return values[i] if i is not None and i < len(values) else None

    rows: list[SheetRow] = []
    for n, values in enumerate(grid[1:], start=2):
        if not any(_s(v) for v in values):
            continue
        raw_code = _s(cell(values, "Employee ID"))
        r = SheetRow(row=n, raw_code=raw_code)
        r.code = normalize_code(cell(values, "Employee ID"))
        if r.code is None:
            r.errors.append(f"no usable employee ID ({raw_code or 'blank'})")
        elif r.code != raw_code:
            r.fixes.append(f"ID {raw_code} → {r.code}")

        r.name = _s(cell(values, "Full name"))
        if not r.name:
            r.errors.append("no name")

        r.raw_department = _s(cell(values, "Department"))
        dept, note = canonical_department(r.raw_department)
        if dept is None:
            r.department = UNASSIGNED
            r.notes.append(f"{note}: stored as {UNASSIGNED}, lowest roles only")
        else:
            r.department = dept
            if dept != r.raw_department:
                r.fixes.append(f"department {r.raw_department} → {dept}")
            if note:
                r.notes.append(note)

        r.designation = _s(cell(values, "Designation"))

        official, fix, err = clean_email(cell(values, "Official email"))
        if err:
            r.notes.append(f"official email unusable ({err}), ignored")
        elif fix:
            r.fixes.append(f"official email: {fix}")
        personal, pfix, perr = clean_email(cell(values, "Personal email"))
        if perr:
            r.notes.append(f"personal email unusable ({perr}), ignored")
        elif pfix:
            r.fixes.append(f"personal email: {pfix}")
        if official and not is_company_email(official):
            if personal and personal != official:
                r.notes.append("official email is not a company address and a personal one is also "
                               "given: the official one is ignored")
            else:
                r.notes.append("official email is not a company address: treated as the personal "
                               "address (needs the code check on first sign-in)")
                personal = official
            official = None
        if personal and personal == official:
            personal = None
        r.official, r.personal = official, personal
        if not official and not personal:
            r.notes.append("no usable email: the account can only sign in with a PIN IT issues")

        status = _s(cell(values, "Status")).lower() or "active"
        if status in ("left", "inactive", "exit"):
            status = "exited"
        if status not in STATUSES:
            r.errors.append(f"status {status!r} is not active or exited")
        r.status = status

        for slug, i in svc_cols.items():
            raw = _s(values[i]) if i < len(values) else ""
            if not raw:
                continue
            key = _role_key(raw, by_slug[slug], role_names.get(slug, {}))
            if key is None:
                r.notes.append(f"{by_slug[slug].name}: {raw!r} is not a role there, cell ignored")
            else:
                r.roles[slug] = key
        rows.append(r)

    _mark_duplicates(rows)
    defaults = _parse_defaults(wb, services, role_names, problems)
    return ParsedSheet(rows=rows, defaults=defaults, problems=problems)


def _mark_duplicates(rows: list[SheetRow]) -> None:
    by_code: dict[str, list[SheetRow]] = defaultdict(list)
    for r in rows:
        if r.code and not r.errors:
            by_code[r.code].append(r)
    for code, group in by_code.items():
        if len(group) < 2:
            continue
        if len({g.identity() for g in group}) == 1:
            for g in group[1:]:
                g.duplicate_of = group[0].row
            continue
        where = ", ".join(str(g.row) for g in group)
        for g in group:
            g.errors.append(f"employee ID {code} is on rows {where} with different details: keep one row")

    # One address typed by several different people is a shared mailbox (a whole lab filling
    # the form from one account), never one person's sign-in: drop it from all of them, say so,
    # and let each sign in with their personal Gmail instead.
    for attr, label in (("official", "official email"), ("personal", "personal email")):
        holders: dict[str, list[SheetRow]] = defaultdict(list)
        for r in rows:
            v = getattr(r, attr)
            if v and not r.errors and not r.duplicate_of:
                holders[v].append(r)
        for v, group in holders.items():
            if len({g.code for g in group}) < 2:
                continue
            where = ", ".join(str(g.row) for g in group)
            for g in group:
                setattr(g, attr, None)
                g.notes.append(f"{label} is shared by rows {where}: looks like a shared mailbox, not used "
                               "as anyone's sign-in" + (" (they sign in with their personal email)" if attr == "official" and g.personal else ""))


def _parse_defaults(wb, services, role_names, problems) -> dict[str, dict[str, str]]:
    if DEFAULTS_SHEET not in wb.sheetnames:
        problems.append(f"no {DEFAULTS_SHEET!r} tab: blank cells get each service's lowest role")
        return {}
    by_slug = {s.slug: s for s in services}
    grid = list(wb[DEFAULTS_SHEET].iter_rows(values_only=True))
    if not grid:
        return {}
    headers = [_s(h) for h in grid[0]]
    _fixed, svc_cols, _ignored = _header_map(headers, services)
    out: dict[str, dict[str, str]] = {}
    for n, values in enumerate(grid[1:], start=2):
        raw = _s(values[0]) if values else ""
        if not raw or raw.lower().startswith("lowest role"):
            continue
        dept, _note = canonical_department(raw)
        if dept is None:
            problems.append(f"{DEFAULTS_SHEET} row {n}: unknown department {raw!r}, ignored")
            continue
        for slug, i in svc_cols.items():
            v = _s(values[i]) if i < len(values) else ""
            if not v:
                continue
            key = _role_key(v, by_slug[slug], role_names.get(slug, {}))
            if key is None or key == NONE:
                problems.append(f"{DEFAULTS_SHEET} row {n}: {v!r} is not a role on {by_slug[slug].name}, ignored")
                continue
            out.setdefault(dept, {})[slug] = key
    return out


def role_names(db: OrmSession) -> dict[str, dict[str, str]]:
    """slug -> {lowercased role name: key}, so a reviewer may type "Manager" for "manager"."""
    out: dict[str, dict[str, str]] = defaultdict(dict)
    for s, r in db.execute(select(Service, ServiceRole).join(ServiceRole, ServiceRole.service_id == Service.id)):
        out[s.slug][r.name.lower()] = r.key
    return out


# ── plan / apply ──────────────────────────────────────────────────────────────
def _find_employee(db: OrmSession, code: str, index: dict[str, Employee]) -> Employee | None:
    return db.scalar(select(Employee).where(Employee.employee_code == code)) or index.get(code)


def run_import(
    db: OrmSession,
    parsed: ParsedSheet,
    services: list[ServiceCol],
    *,
    actor: User | None,
    skip: set[str] | None = None,
) -> dict:
    """Apply the sheet to the session. The caller commits, or rolls back for a dry run."""
    now = datetime.now(timezone.utc)
    skip = {k.lower() for k in (skip or set())}
    index = {normalize_code(e.employee_code): e for e in db.scalars(select(Employee)) if normalize_code(e.employee_code)}
    out_rows: list[dict] = []
    sheet_users: dict[str, tuple[User, Employee, SheetRow]] = {}
    counts = Counter()

    for r in parsed.rows:
        entry = {"row": r.row, "employee_code": r.code or r.raw_code, "name": r.name,
                 "department": r.department, "errors": list(r.errors), "fixes": list(r.fixes),
                 "notes": list(r.notes), "changes": []}
        out_rows.append(entry)
        if r.errors:
            entry["status"] = "reject"
            counts["rejected"] += 1
            continue
        if r.duplicate_of:
            entry["status"] = "duplicate"
            entry["notes"].append(f"same person as row {r.duplicate_of}, listed twice: row {r.duplicate_of} is used")
            counts["duplicates"] += 1
            continue

        emp = _find_employee(db, r.code, index)
        if emp is not None and emp.job_title == FUNCTIONAL_JOB_TITLE:
            entry["status"] = "reject"
            entry["errors"].append(f"{emp.employee_code} is a shared mailbox account, not a person")
            counts["rejected"] += 1
            continue
        if r.official:
            clash = db.scalar(select(Employee).where(func.lower(Employee.work_email) == r.official))
            if clash is not None and (emp is None or clash.id != emp.id):
                entry["status"] = "reject"
                entry["errors"].append(f"official email already belongs to employee {clash.employee_code}")
                counts["rejected"] += 1
                continue
            uclash = db.scalar(select(User).where(func.lower(User.login_email) == r.official))
            if uclash is not None and (emp is None or uclash.employee_id != emp.id):
                entry["status"] = "reject"
                entry["errors"].append("official email is already another account's sign-in address")
                counts["rejected"] += 1
                continue

        changes = entry["changes"]
        if emp is None:
            emp = Employee(
                employee_code=r.code, full_name=r.name, work_email=r.official,
                hr_department=r.department, division=r.department, job_title=r.designation,
                band="N/A", notes="people sheet",
            )
            db.add(emp)
            db.flush()
            index[r.code] = emp
            status = "create"
        else:
            if emp.employee_code != r.code:
                entry["notes"].append(f"already stored as {emp.employee_code}; code kept as it is")
            for label, attr, new in (("name", "full_name", r.name), ("department", "hr_department", r.department),
                                     ("designation", "job_title", r.designation), ("official email", "work_email", r.official)):
                if new and (getattr(emp, attr) or "") != new:
                    changes.append({"field": label, "from": getattr(emp, attr), "to": new})
                    setattr(emp, attr, new)
                    if attr == "hr_department":
                        emp.division = new
            status = "update" if changes else "unchanged"

        user = db.scalar(select(User).where(User.employee_id == emp.id))
        if user is None:
            if r.official:
                user = User(employee_id=emp.id, auth_type="google", login_email=r.official, is_active=r.status == "active")
            else:
                # No company address: a PIN account whose PIN nobody knows. The person gets in
                # through their personal Gmail + employee code (onboarding sets the real PIN),
                # or IT issues a PIN.
                user = User(employee_id=emp.id, auth_type="local_pin", pin_hash=hash_pin(f"{secrets.randbelow(10**8):08d}"),
                            is_active=r.status == "active")
            db.add(user)
            db.flush()
            if status == "unchanged":
                status = "update"
            changes.append({"field": "account", "from": None, "to": "created"})
        else:
            if r.official and not user.login_email:
                user.login_email, user.auth_type = r.official, "google"
                changes.append({"field": "Google sign-in", "from": None, "to": r.official})
            elif r.official and user.login_email and user.login_email.lower() != r.official:
                entry["notes"].append("signs in with a different address already; sign-in address not changed")
            if r.status == "exited" and user.is_active:
                user.is_active = False
                for s in db.scalars(select(Session).where(Session.user_id == user.id, Session.revoked_at.is_(None))):
                    s.revoked_at = now
                changes.append({"field": "account", "from": "active", "to": "switched off (exited)"})
            elif r.status == "active" and not user.is_active:
                entry["notes"].append("account is switched off in MM OS; the sheet does not switch it back on")
            if changes and status == "unchanged":
                status = "update"

        if r.personal:
            other = db.get(PersonalEmail, r.personal)
            owner_login = db.scalar(select(User).where(func.lower(User.login_email) == r.personal))
            current = personal_email_for(db, user)
            if (other is not None and other.user_id != user.id) or (owner_login is not None and owner_login.id != user.id):
                entry["notes"].append("personal email is already used by another account, ignored")
            elif current is None or current.email != r.personal:
                if current is not None:
                    db.delete(current)
                    db.flush()
                db.add(PersonalEmail(email=r.personal, user_id=user.id))
                changes.append({"field": "personal email", "from": current.email if current else None, "to": r.personal})
                if status == "unchanged":
                    status = "update"
        db.flush()
        entry["status"] = status
        counts[status] += 1
        if r.status == "active":  # an exited person is switched off, never given access
            sheet_users[emp.employee_code] = (user, emp, r)

    grants = _apply_grants(db, parsed, services, sheet_users, actor=actor, skip=skip, now=now)
    for g in grants:
        counts[f"grants_{g['kind']}"] += 1

    dept_map = Counter((r.raw_department, r.department) for r in parsed.rows if r.raw_department != r.department)
    return {
        "summary": {"rows": len(parsed.rows), **counts},
        "rows": out_rows,
        "grants": grants,
        "departments": [{"from": a, "to": b, "rows": n} for (a, b), n in sorted(dept_map.items(), key=lambda x: (x[0][1], x[0][0]))],
        "defaults": parsed.defaults,
        "lowest_roles": {s.slug: s.lowest for s in services},
        "warnings": parsed.problems + [f"{s.name}: no lowest role is set, so a blank cell with no department "
                                       "default gives no access there" for s in services if not s.lowest],
    }


def _apply_grants(db, parsed, services, sheet_users, *, actor, skip, now) -> list[dict]:
    out: list[dict] = []
    only = {u.id for u, _e, _r in sheet_users.values()}
    if not only:
        return out
    by_slug = {s.slug: s for s in db.scalars(select(Service).where(Service.slug.in_([c.slug for c in services])))}
    for col in services:
        service = by_slug.get(col.slug)
        if service is None:
            continue
        explicit = {code: r.roles[col.slug] for code, (_u, _e, r) in sheet_users.items()
                    if r.roles.get(col.slug) and r.roles[col.slug] != NONE}
        rules = [{"department": [dept], "role": m[col.slug]} for dept, m in parsed.defaults.items() if m.get(col.slug)]
        assign = {"people": explicit, "rules": rules, "default": col.lowest}
        plan = roles_io.RolePlan(service=col.slug)
        # `none` means no access here, so those people never get a default or lowest role either.
        said_none = {u.id for u, _e, r in sheet_users.values() if r.roles.get(col.slug) == NONE}
        roles_io.assign_roles(db, service, assign, plan, actor=actor, mode="missing", people_override=True,
                              only_user_ids=only - said_none, skip=skip, reason="people sheet")
        for g in plan.grants_created:
            out.append({"kind": "created", "key": f"{g['employee_code']}:{col.slug}".lower(), "employee_code": g["employee_code"],
                        "name": g["name"], "service": col.slug, "service_name": col.name, "from": None, "to": g["role"],
                        "source": g["source"], "direction": "up"})
        for g in plan.grants_changed:
            out.append({"kind": "changed", "key": f"{g['employee_code']}:{col.slug}".lower(), "employee_code": g["employee_code"],
                        "name": g["name"], "service": col.slug, "service_name": col.name, "from": g["from"], "to": g["to"],
                        "source": g["source"], "direction": g["direction"]})

        role_by_id = {r.id: r.key for r in service.roles}
        for code, (user, emp, r) in sheet_users.items():
            if r.roles.get(col.slug) != NONE:
                continue
            key = f"{emp.employee_code}:{col.slug}".lower()
            grant = db.scalar(select(Grant).where(Grant.user_id == user.id, Grant.service_id == service.id))
            if grant is None or key in skip:
                continue
            out.append({"kind": "removed", "key": key, "employee_code": emp.employee_code, "name": emp.full_name,
                        "service": col.slug, "service_name": col.name, "from": role_by_id.get(grant.service_role_id),
                        "to": None, "source": "people", "direction": "down"})
            db.add(Revocation(subject=user.subject, service_id=service.id, reason="grant_removed",
                              revoked_by=actor.id if actor else None, revoked_at=now,
                              purge_after=now + roles_io._REVOCATION_TTL))
            db.delete(grant)
        db.flush()
    return out
