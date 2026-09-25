"""People sheet: template, clean-up, dry run and apply (app/people_sheet.py). Fake people only."""
from __future__ import annotations

import io
import json

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app import models
from app.departments import canonical_department, clean_email, normalize_code
from app.onboarding import PersonalEmail
from app.people_sheet import FIXED, build_workbook, service_columns


# ── clean-up rules ──────────────────────────────────────────────────────────────
@pytest.mark.parametrize("raw, code", [
    ("MM00115", "MM115"), ("mm-115", "MM115"), ("MM 115", "MM115"), (115.0, "MM115"), ("115", "MM115"),
    ("MM05", "MM5"), ("NA", None), ("..", None), ("test1234", None), ("MM0", None), (None, None),
])
def test_normalize_code(raw, code):
    assert normalize_code(raw) == code


@pytest.mark.parametrize("raw, dept", [
    ("P-HUB ", "P-Hub"), ("Phub", "P-Hub"), ("Production", "P-Hub"), ("SPOKE", "P-Spoke"),
    ("ENGINEERING", "P-Spoke"), ("MAINTENANCE ", "P-Spoke"), ("2nd life ", "P-Spoke"),
    ("STORE", "Stores"), ("Material Management", "Stores"), ("LOGISTICS ", "Logistics"),
    ("Business Development ", "BD & Operations"), ("BD/OPERATIONS", "BD & Operations"),
    ("Business Operations", "BD & Operations"), ("FINANCE OPERATIONS", "Finance"),
    ("Human Resource ", "HR"), ("HR & Stores", "HR"), ("Research and Development", "R&D"),
    ("QA/QC", "QA/QC"), ("it", "IT"), ("IT ADMIN ", "IT"), ("CTO", "CXO Office"), ("Corporate", "CXO Office"),
    ("STRATEGY AND OPERATIONS ", "StratOps"), ("EHS", "EHS"), ("PROJECTS", "Projects"),
])
def test_department_aliases(raw, dept):
    assert canonical_department(raw)[0] == dept


def test_unknown_department_is_reported_not_guessed():
    dept, note = canonical_department("Space Program")
    assert dept is None and "unknown department" in note


@pytest.mark.parametrize("raw, email, fixed", [
    ("a.b@gamail.com", "a.b@gmail.com", True), ("a.b@gmail .com", "a.b@gmail.com", True),
    ("a.b@gmail com", "a.b@gmail.com", True), (" A.B@M-Mines.com ", "a.b@m-mines.com", False),
])
def test_clean_email(raw, email, fixed):
    got, fix, err = clean_email(raw)
    assert got == email and bool(fix) == fixed and err is None


def test_clean_email_rejects_two_addresses_and_junk():
    assert clean_email("x@m-mines.com , y@mimines.com")[2]
    assert clean_email("not an email")[2]


# ── fixtures ────────────────────────────────────────────────────────────────────
@pytest.fixture()
def admin(make_user, sign_in):
    user = make_user(is_platform_admin=True)
    sign_in(user)
    return user


@pytest.fixture()
def services(db, make_service):
    ic, ic_roles = make_service(slug="itemcode", name="Item Code Studio", roles=("public", "associate", "manager", "admin"))
    for i, r in enumerate(ic_roles.values()):
        r.sort_order = (i + 1) * 10
    ic_roles["public"].is_default = True
    sd, sd_roles = make_service(slug="servicedesk", name="Service Desk", roles=("requester", "agent", "admin"))
    for i, r in enumerate(sd_roles.values()):
        r.sort_order = (i + 1) * 10
    sd_roles["requester"].is_default = True
    make_service(slug="erpnext", name="ERPNext", roles=("user",), launch_mode="external")
    db.commit()
    return {"itemcode": (ic, ic_roles), "servicedesk": (sd, sd_roles)}


def _row(code, name, dept, official="", personal="", status="", itemcode="", servicedesk="", designation="Engineer"):
    return {"Employee ID": code, "Full name": name, "Department": dept, "Designation": designation,
            "Official email": official, "Personal email": personal, "Status": status,
            "roles": {k: v for k, v in (("itemcode", itemcode), ("servicedesk", servicedesk)) if v}}


def _workbook(db, people, defaults=None) -> bytes:
    return build_workbook(service_columns(db), people=people, defaults=defaults)


def _upload(client, data: bytes, *, dry_run=True, skip=None):
    files = {"file": ("people.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    form = {"skip": json.dumps(skip)} if skip else {}
    return client.post(f"/api/admin/people/import?dry_run={'true' if dry_run else 'false'}", files=files, data=form)


# ── template ────────────────────────────────────────────────────────────────────
def test_template_has_tabs_dropdowns_and_no_external_services(client, admin, services):
    resp = client.get("/api/admin/people/template.xlsx")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert {"People", "Department defaults", "Instructions", "Lists"} <= set(wb.sheetnames)
    headers = [c.value for c in wb["People"][1]]
    assert headers[: len(FIXED)] == FIXED
    assert "Item Code Studio" in headers and "Service Desk" in headers and "ERPNext" not in headers
    assert len(wb["People"].data_validations.dataValidation) >= 4
    lists = [c.value for c in wb["Lists"]["C"]]
    assert lists[:6] == ["Item Code Studio", "public", "associate", "manager", "admin", "none"]


def test_template_requires_admin(client, make_user, sign_in):
    sign_in(make_user())
    assert client.get("/api/admin/people/template.xlsx").status_code == 403


# ── dry run ─────────────────────────────────────────────────────────────────────
def test_dry_run_reports_fixes_rejects_and_duplicates_and_writes_nothing(client, admin, services, db):
    people = [
        _row("MM00115", "Fake One", "MAINTENANCE", "fake.one@m-mines.com", "fake.one@gamail.com"),
        _row("MM-115", "Fake One", "ENGINEERING", "fake.one@m-mines.com", "fake.one@gamail.com"),  # same person twice
        _row("NA", "No Id", "Finance", "no.id@m-mines.com"),
        _row("MM200", "Clash A", "Purchase", "clash.a@m-mines.com"),
        _row("MM200", "Clash B", "Stores", "clash.b@m-mines.com"),
        _row("MM300", "Gmail Only", "P-HUB", "gmail.only@gmail.com"),
        _row("MM400", "Odd Dept", "Space Program", "odd.dept@m-mines.com"),
    ]
    before = db.scalar(select(models.Employee).where(models.Employee.employee_code == "MM115"))
    resp = _upload(client, _workbook(db, people))
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["dry_run"] is True
    rows = {r["row"]: r for r in body["rows"]}
    assert rows[2]["status"] == "create"
    assert any("MM00115 → MM115" in f for f in rows[2]["fixes"])
    assert any("gmail.com" in f for f in rows[2]["fixes"])
    assert rows[3]["status"] == "duplicate"
    assert rows[4]["status"] == "reject" and "no usable employee ID" in rows[4]["errors"][0]
    assert rows[5]["status"] == rows[6]["status"] == "reject"
    assert rows[7]["status"] == "create" and any("not a company address" in n for n in rows[7]["notes"])
    assert rows[8]["department"] == "Unassigned"
    assert {"from": "MAINTENANCE", "to": "P-Spoke", "rows": 1} in body["departments"]
    assert before is None
    assert db.scalar(select(models.Employee).where(models.Employee.employee_code == "MM115")) is None


# ── apply ───────────────────────────────────────────────────────────────────────
def test_apply_creates_people_accounts_and_access_with_the_right_precedence(client, admin, services, db):
    defaults = {"Purchase": {"itemcode": "associate"}}
    people = [
        _row("MM1", "Buyer Explicit", "Purchase", "buyer.one@m-mines.com", itemcode="manager"),
        _row("MM2", "Buyer Default", "Purchase", "buyer.two@m-mines.com", "buyer.two@gmail.com"),
        _row("MM3", "Finance Lowest", "Finance", "fin.three@m-mines.com", servicedesk="agent"),
        _row("MM4", "Gmail Only", "Stores", personal="gmail.four@gmail.com"),
    ]
    resp = _upload(client, _workbook(db, people, defaults), dry_run=False)
    assert resp.status_code == 200, resp.text
    s = resp.json()["summary"]
    assert s["create"] == 4 and s.get("rejected", 0) == 0

    def role(code, slug):
        emp = db.scalar(select(models.Employee).where(models.Employee.employee_code == code))
        user = db.scalar(select(models.User).where(models.User.employee_id == emp.id))
        g = db.scalar(select(models.Grant).join(models.Service).where(models.Grant.user_id == user.id, models.Service.slug == slug))
        return db.get(models.ServiceRole, g.service_role_id).key if g else None

    assert role("MM1", "itemcode") == "manager"       # filled-in cell
    assert role("MM2", "itemcode") == "associate"     # department default
    assert role("MM3", "itemcode") == "public"        # lowest role
    assert role("MM3", "servicedesk") == "agent"
    assert role("MM1", "erpnext") is None              # external: never touched

    u2 = db.scalar(select(models.User).where(models.User.login_email == "buyer.two@m-mines.com"))
    assert u2.auth_type == "google" and u2.pin_set_at is None
    pe = db.get(PersonalEmail, "buyer.two@gmail.com")
    assert pe.user_id == u2.id and pe.verified_at is None
    emp4 = db.scalar(select(models.Employee).where(models.Employee.employee_code == "MM4"))
    u4 = db.scalar(select(models.User).where(models.User.employee_id == emp4.id))
    assert u4.auth_type == "local_pin" and u4.login_email is None and db.get(PersonalEmail, "gmail.four@gmail.com")


def test_reupload_cells_replace_blanks_fill_gaps_none_removes_and_skip_is_respected(client, admin, services, db, make_employee, make_user, make_grant):
    ic, ic_roles = services["itemcode"]
    sd, sd_roles = services["servicedesk"]
    emp = make_employee(employee_code="MM10", hr_department="Purchase", work_email="hand.set@m-mines.com")
    user = make_user(employee=emp)
    make_grant(user, ic, ic_roles["admin"])        # hand-set; a blank cell must keep it
    make_grant(user, sd, sd_roles["agent"])        # the sheet says none
    emp2 = make_employee(employee_code="MM11", hr_department="Purchase", work_email="promote.me@m-mines.com")
    user2 = make_user(employee=emp2)
    make_grant(user2, ic, ic_roles["associate"])

    people = [
        _row("MM10", emp.full_name, "Purchase", "hand.set@m-mines.com", servicedesk="none"),
        _row("MM11", emp2.full_name, "Purchase", "promote.me@m-mines.com", itemcode="manager", servicedesk="agent"),
    ]
    data = _workbook(db, people, {"Purchase": {"itemcode": "associate"}})
    preview = _upload(client, data).json()
    kinds = {(g["employee_code"], g["service"]): (g["kind"], g["from"], g["to"], g["direction"]) for g in preview["grants"]}
    assert ("MM10", "itemcode") not in kinds                                        # blank keeps admin
    assert kinds[("MM10", "servicedesk")] == ("removed", "agent", None, "down")
    assert kinds[("MM11", "itemcode")] == ("changed", "associate", "manager", "up")
    assert kinds[("MM11", "servicedesk")][0] == "created"

    resp = _upload(client, data, dry_run=False, skip=["mm11:servicedesk"])
    assert resp.status_code == 200
    db.expire_all()
    held = {(g.user_id, g.service_id): db.get(models.ServiceRole, g.service_role_id).key for g in db.scalars(select(models.Grant))}
    assert held[(user.id, ic.id)] == "admin"
    assert (user.id, sd.id) not in held
    assert held[(user2.id, ic.id)] == "manager"
    assert (user2.id, sd.id) not in held  # unticked in the dry run
    assert db.scalar(select(models.Revocation).where(models.Revocation.subject == user.subject)) is not None


def test_exited_switches_the_account_off(client, admin, services, db, make_employee, make_user):
    emp = make_employee(employee_code="MM20", hr_department="Finance", work_email="leaver@m-mines.com")
    user = make_user(employee=emp)
    resp = _upload(client, _workbook(db, [_row("MM20", emp.full_name, "Finance", "leaver@m-mines.com", status="exited")]), dry_run=False)
    assert resp.status_code == 200
    db.refresh(user)
    assert user.is_active is False


def test_shared_mailbox_ids_are_refused(client, admin, services, db, make_employee, make_user):
    from app.provision import FUNCTIONAL_JOB_TITLE
    box = make_employee(employee_code="MM900", job_title=FUNCTIONAL_JOB_TITLE, work_email="purchase.box@m-mines.com")
    make_user(employee=box)
    body = _upload(client, _workbook(db, [_row("MM900", "Some Person", "Purchase", "some.person@m-mines.com")])).json()
    assert body["rows"][0]["status"] == "reject"


def test_not_a_workbook_is_a_clean_422(client, admin, services):
    resp = client.post("/api/admin/people/import", files={"file": ("x.xlsx", b"not a zip", "application/octet-stream")})
    assert resp.status_code == 422 and resp.json()["error"] == "bad_workbook"


def test_none_on_a_new_person_gives_no_lowest_role_and_reports_nothing_removed(client, admin, services, db):
    body = _upload(client, _workbook(db, [_row("MM30", "No Desk", "Finance", "no.desk@m-mines.com", servicedesk="none")])).json()
    keys = {(g["service"], g["kind"]) for g in body["grants"]}
    assert ("servicedesk", "created") not in keys and ("servicedesk", "removed") not in keys
    assert ("itemcode", "created") in keys


def test_one_company_address_typed_by_several_people_is_a_shared_mailbox_not_a_sign_in(client, admin, services, db):
    people = [
        _row("MM41", "Lab One", "QA/QC", "lab.box@m-mines.com", "lab.one@gmail.com"),
        _row("MM42", "Lab Two", "QA/QC", "lab.box@m-mines.com", "lab.two@gmail.com"),
    ]
    body = _upload(client, _workbook(db, people), dry_run=False).json()
    assert [r["status"] for r in body["rows"]] == ["create", "create"]
    assert all(any("shared mailbox" in n for n in r["notes"]) for r in body["rows"])
    assert db.scalar(select(models.User).where(models.User.login_email == "lab.box@m-mines.com")) is None
    assert db.get(PersonalEmail, "lab.one@gmail.com") is not None
